#!/usr/bin/env python3
"""PinchBench 评测报告汇总脚本（中文整合版）。

将多个模型的评测结果 JSON 按各维度分析，整合到一个中文 Excel 文件中。
等价整合 analysis.py（结果维度分析）与 export_tasks_to_excel.py（用例信息），
并将用例名称、输入、预期行为、评分标准等以中文展示（来源 tasks/cn/*.md）。

Sheet 结构（共 7 + N，N=模型数）：
  1. 总览
  2. 用例对比明细
  3. Agent能力对比
  4. 分类对比
  5. 场景对比
  6. 难度等级对比
  7. 模型分差矩阵
  8. 评分详情_{模型}（每个模型一个）

用法:
  python generate_eval_report.py -d <结果目录> [<结果目录> ...] [-o <输出目录>]
"""

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("需要 openpyxl 库，请执行: pip install openpyxl")

try:
    import yaml
except ImportError:
    sys.exit("需要 pyyaml 库，请执行: pip install pyyaml")


# --------------------------------------------------------------------------- #
# 常量：样式、命名规则
# --------------------------------------------------------------------------- #
# 形如 0004_xopdeepseekv4pro.json：0 开头 + 下划线 + 模型 id + .json
RESULT_FILE_PATTERN = re.compile(r"^0\w*_.+\.json$")

HEADER_FONT_WHITE = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
CENTER = Alignment(horizontal="center")
WRAP_TOP = Alignment(wrap_text=True, vertical="top")

# Excel 单元格字符上限（实际 32767，留余量）
CELL_MAX_LEN = 32000

# Agent能力强项/短板排名时，能力涉及用例数的最小阈值（过滤小样本，避免偶然高/低分干扰）
CAP_RANK_MIN_COUNT = 5


# --------------------------------------------------------------------------- #
# 常量：中英映射（来源 docs/reference/*.md）
# --------------------------------------------------------------------------- #
# 11 个分类，按 manifest.yaml 顺序（来源 docs/reference/category.md）
CATEGORY_ORDER = [
    "productivity", "research", "writing", "coding", "analysis",
    "csv_analysis", "log_analysis", "meeting_analysis", "memory",
    "skills", "integrations",
]
CATEGORY_ZH = {
    "productivity": "生产力",
    "research": "调研",
    "writing": "写作",
    "coding": "编程",
    "analysis": "综合分析",
    "csv_analysis": "CSV 数据分析",
    "log_analysis": "日志分析",
    "meeting_analysis": "会议分析",
    "memory": "记忆",
    "skills": "技能",
    "integrations": "集成",
}

# 8 大场景展示顺序与中文名（来源 docs/reference/scene.md）
SCENE_ORDER = [
    "finance_investment_research", "deep_research_report", "science_tech_medical_qa",
    "data_retrieval_analysis", "content_creation_multimedia", "enterprise_product_intel",
    "skill_lifecycle", "local_env_scripting",
]
SCENE_ZH = {
    "finance_investment_research": "金融投研与企业价值评估",
    "deep_research_report": "深度搜索与专题研究报告",
    "science_tech_medical_qa": "科学技术、医学与计算问答",
    "data_retrieval_analysis": "数据库检索、表格整理与数据分析",
    "content_creation_multimedia": "内容创作、PPT、网页与多媒体生成",
    "enterprise_product_intel": "企业产品情报与业务信息助手",
    "skill_lifecycle": "Skill发现、创建、安装与调用",
    "local_env_scripting": "本地环境、命令执行与脚本任务",
}
# scene -> S1~S8 分组编号
SCENE_GROUP = {s: f"S{i + 1}" for i, s in enumerate(SCENE_ORDER)}

# 20 个标准能力标签顺序与中文名（来源 docs/reference/agent-capability-dimensions.md）
CAPABILITY_ORDER = [
    "instruction_following", "context_memory", "output_format", "hallucination_resistance",
    "tool_usage", "multimodal_perception", "data_extraction", "information_retrieval",
    "multi_step_reasoning", "planning", "domain_reasoning", "code_generation",
    "service_integration", "text_generation", "self_correction", "uncertainty_handling",
    "safety_awareness", "concurrency_management", "multi_agent", "adaptive_learning",
]
CAPABILITY_ZH = {
    "instruction_following": "指令遵循与约束理解",
    "context_memory": "上下文记忆与状态管理",
    "output_format": "输出格式适配",
    "hallucination_resistance": "幻觉抑制",
    "tool_usage": "工具调用",
    "multimodal_perception": "多模态感知",
    "data_extraction": "数据提取与处理",
    "information_retrieval": "信息检索与综合",
    "multi_step_reasoning": "多步推理",
    "planning": "规划与任务分解",
    "domain_reasoning": "领域推理",
    "code_generation": "代码生成与理解",
    "service_integration": "外部服务集成",
    "text_generation": "自然语言生成",
    "self_correction": "自我纠错与反思",
    "uncertainty_handling": "不确定性处理",
    "safety_awareness": "安全与权限意识",
    "concurrency_management": "并发与优先级管理",
    "multi_agent": "多Agent协作",
    "adaptive_learning": "自适应学习",
}

# 难度等级展示顺序与中文名
DIFFICULTY_ORDER = ["L1", "L2", "L3", "L4"]
DIFFICULTY_ZH = {
    "L1": "L1 入门",
    "L2": "L2 进阶",
    "L3": "L3 困难",
    "L4": "L4 专家",
}


def bilingual(zh: str, en: str) -> str:
    """中文(英文) 双标格式。zh 为空时仅返回 en。"""
    if not en:
        return zh or ""
    if not zh or zh == en:
        return en
    return f"{zh}({en})"


# --------------------------------------------------------------------------- #
# 项目根 / 用例元数据加载
# --------------------------------------------------------------------------- #
class TaskMeta:
    """单个用例的中文元数据（来源 tasks/cn/*.md，回退 tasks/*.md）。"""

    def __init__(self):
        self.name_zh = ""
        self.category_en = ""
        self.scene_en = ""
        self.sub_scene_zh = ""
        self.difficulty = ""
        self.capabilities_en: list[str] = []
        self.prompt = ""          # 中文输入
        self.expected = ""        # 中文预期行为
        self.criteria = ""        # 中文评分标准（带编号）


def find_project_root(start: Path) -> Path | None:
    """向上查找 PinchBench 项目根：同时存在 scripts/lib_grading.py 与 tasks/ 目录。"""
    cur = start.resolve()
    for cand in [cur, *cur.parents]:
        if (cand / "scripts" / "lib_grading.py").exists() and (cand / "tasks").is_dir():
            return cand
    return None


def _parse_frontmatter_and_sections(text: str) -> tuple[dict, dict]:
    """解析 md：返回 (frontmatter dict, sections dict)。"""
    fm: dict = {}
    body = text
    m = re.match(r"^---\s*\n(.*?)\n---\s*\n(.*)$", text, re.DOTALL)
    if m:
        try:
            parsed = yaml.safe_load(m.group(1))
            if isinstance(parsed, dict):
                fm = parsed
        except yaml.YAMLError:
            fm = {}
        body = m.group(2)

    sections: dict[str, str] = {}
    current = None
    buf: list[str] = []
    for line in body.split("\n"):
        h = re.match(r"^##\s+(.+)$", line)
        if h:
            if current:
                sections[current] = "\n".join(buf).strip()
            current = h.group(1).strip()
            buf = []
        elif current:
            buf.append(line)
    if current:
        sections[current] = "\n".join(buf).strip()
    return fm, sections


def _extract_criteria(text: str) -> str:
    """从 ## Grading Criteria 的复选框列表提取为带编号的字符串。"""
    items = []
    for line in text.split("\n"):
        m = re.match(r"^-\s+\[[ x]\]\s+(.+)$", line.strip())
        if m:
            items.append(m.group(1))
    return "\n".join(f"{i + 1}. {it}" for i, it in enumerate(items))


def load_task_meta(project_root: Path | None) -> dict[str, TaskMeta]:
    """加载所有用例的中文元数据。

    中文字段（name/sub_scene/prompt/expected/criteria）取自 tasks/cn/*.md，
    英文 enum（category/scene/capabilities）取自 tasks/*.md（用于双标）。
    difficulty 两处一致，优先取英文目录。
    """
    metas: dict[str, TaskMeta] = {}
    if project_root is None:
        return metas

    en_dir = project_root / "tasks"
    cn_dir = project_root / "tasks" / "cn"

    for en_file in sorted(en_dir.glob("task_*.md")):
        task_id = en_file.stem
        if task_id == "task_XX_name":
            continue
        try:
            en_fm, _ = _parse_frontmatter_and_sections(en_file.read_text(encoding="utf-8"))
        except OSError:
            continue

        meta = TaskMeta()
        meta.category_en = en_fm.get("category", "") or ""
        meta.scene_en = en_fm.get("scene", "") or ""
        meta.difficulty = en_fm.get("difficulty", "") or ""
        caps = en_fm.get("capabilities") or []
        meta.capabilities_en = caps if isinstance(caps, list) else []

        cn_file = cn_dir / f"{task_id}.md"
        if cn_file.exists():
            try:
                cn_fm, cn_sec = _parse_frontmatter_and_sections(
                    cn_file.read_text(encoding="utf-8")
                )
                meta.name_zh = cn_fm.get("name", "") or ""
                meta.sub_scene_zh = cn_fm.get("sub_scene", "") or ""
                meta.prompt = (cn_sec.get("Prompt", "") or "").strip()
                meta.expected = (cn_sec.get("Expected Behavior", "") or "").strip()
                meta.criteria = _extract_criteria(cn_sec.get("Grading Criteria", "") or "")
            except OSError:
                pass
        else:
            print(f"警告：缺少中文用例文件，将回退英文名: {cn_file}", file=sys.stderr)
            meta.name_zh = en_fm.get("name", "") or ""

        metas[task_id] = meta

    return metas


# --------------------------------------------------------------------------- #
# 结果文件加载与去重
# --------------------------------------------------------------------------- #
def find_result_files(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*.json") if RESULT_FILE_PATTERN.match(p.name))


def infer_result_root(files: list[Path]) -> Path:
    """从结果文件推断评测结果根目录。

    结果文件通常位于 <结果根>/<模型目录>/0xxx_<model>.json，
    取各文件父目录的共同上层作为结果根；只有单个模型目录时取其父目录。
    """
    parents = {f.parent.resolve() for f in files}
    if len(parents) == 1:
        # 全部在同一模型目录下 → 结果根是它的父目录
        return next(iter(parents)).parent
    # 多个模型目录 → 取共同父目录
    import os
    common = Path(os.path.commonpath([str(p) for p in parents]))
    return common


def dedup_tasks(tasks: list[dict]) -> list[dict]:
    """按 task_id 去重，保留首条。

    结果 JSON 中每个 task_id 会按运行次数重复出现，但 grading.mean 已是跨轮聚合值，
    各条相同。聚合维度统计必须先去重，否则 task_count 被放大。
    """
    seen: set[str] = set()
    out: list[dict] = []
    for t in tasks:
        tid = t.get("task_id")
        if tid and tid not in seen:
            seen.add(tid)
            out.append(t)
    return out


class ModelResult:
    """单个模型的评测结果（已去重 tasks）。"""

    def __init__(self, path: Path, data: dict):
        self.path = path
        self.data = data
        self.model = data.get("model", path.stem)
        self.run_id = data.get("run_id", "")
        self.suite = data.get("suite", "")
        self.timestamp = data.get("timestamp", "")
        self.efficiency = data.get("efficiency", {}) or {}
        self.category_scores = data.get("category_scores", {}) or {}
        self.tasks = dedup_tasks(data.get("tasks", []))
        # task_id -> mean 得分
        self.task_score: dict[str, float] = {
            t["task_id"]: float(t.get("grading", {}).get("mean", 0.0))
            for t in self.tasks if t.get("task_id")
        }

    @property
    def total_pct(self) -> float:
        """总分率（%）：基于去重后 task 的 mean 平均，每 task 满分 1.0。"""
        if not self.task_score:
            return 0.0
        return sum(self.task_score.values()) / len(self.task_score) * 100


# --------------------------------------------------------------------------- #
# 维度聚合（基于 metas + 各模型 task_score）
# --------------------------------------------------------------------------- #
def avg_pct(model: ModelResult, task_ids: list[str]) -> float | None:
    """某模型在一组用例上的平均得分率（%）。无交集返回 None。"""
    vals = [model.task_score[t] for t in task_ids if t in model.task_score]
    if not vals:
        return None
    return sum(vals) / len(vals) * 100


def fmt_pct(v: float | None) -> str:
    return f"{v:.1f}%" if v is not None else "-"


# --------------------------------------------------------------------------- #
# transcript 读取
# --------------------------------------------------------------------------- #
def read_transcript_raw(result_path: Path, run_id: str, task_id: str) -> str:
    """读取 {run_id}_transcripts/{task_id}.jsonl 全文（不截断）。

    找不到返回空串。供摘要解析使用——必须用未截断文本，
    否则超长首行会被切断导致 JSONL 解析失败。
    """
    base = result_path.parent
    candidates = []
    if run_id:
        candidates.append(base / f"{run_id}_transcripts" / f"{task_id}.jsonl")
    candidates += list(base.glob(f"*_transcripts/{task_id}.jsonl"))
    for c in candidates:
        if c.exists():
            try:
                return c.read_text(encoding="utf-8", errors="replace")
            except OSError:
                continue
    return ""


def read_transcript(result_path: Path, run_id: str, task_id: str) -> str:
    """读取 transcript 全文，超长截断（供 Excel「实际结果」列）。"""
    text = read_transcript_raw(result_path, run_id, task_id)
    if not text:
        return "无 transcript 记录"
    if len(text) > CELL_MAX_LEN:
        text = text[:CELL_MAX_LEN] + "\n...（已截断）"
    return text


def format_breakdown(grading: dict) -> str:
    """组装多轮检查点得分明细。

    格式：
      第1轮 (score=1.0):
        key: val
      第2轮 (score=0.83):
        ...
    """
    runs = grading.get("runs", []) or []
    blocks = []
    for i, run in enumerate(runs, 1):
        score = run.get("score", 0.0)
        bd = run.get("breakdown", {}) or {}
        lines = [f"第{i}轮 (score={score}):"]
        for k, v in bd.items():
            lines.append(f"  {k}: {v}")
        blocks.append("\n".join(lines))
    return "\n".join(blocks)


def format_lost_points(grading: dict) -> str:
    """组装失分点（纯规则提取，按轮次分别列）。

    每轮包含两部分：
      - 失分检查点：breakdown 中得分 < 1.0 的项（key=得分），区分完全失分与部分失分
      - 裁判判词：该轮 grading_type 为 llm_judge/hybrid 且 notes 非空时附上

    满分用例（各轮均无失分）整体返回「无失分」。
    """
    runs = grading.get("runs", []) or []
    blocks = []
    any_loss = False

    for i, run in enumerate(runs, 1):
        score = run.get("score", 0.0)
        bd = run.get("breakdown", {}) or {}
        gtype = run.get("grading_type", "")
        notes = (run.get("notes") or "").strip()

        zero = [k for k, v in bd.items() if v == 0.0]
        partial = [(k, v) for k, v in bd.items() if 0.0 < v < 1.0]

        lines = [f"第{i}轮 (score={score}):"]
        if zero:
            lines.append("  完全失分: " + ", ".join(zero))
        if partial:
            lines.append("  部分失分: " + ", ".join(f"{k}={v}" for k, v in partial))
        if not zero and not partial:
            lines.append("  本轮无失分")
        else:
            any_loss = True
        # 裁判判词（仅 llm_judge / hybrid）
        if gtype in ("llm_judge", "hybrid") and notes:
            lines.append(f"  裁判: {notes}")
        blocks.append("\n".join(lines))

    if not any_loss:
        return "无失分"
    return "\n".join(blocks)


def summarize_transcript(text: str, max_len: int = 6000) -> str:
    """为 LLM 分析压缩 transcript：保留 assistant 文本与 toolCall/toolResult 关键片段。

    transcript 是 JSONL（每行一个事件）。优先抽取与模型行为相关的事件，
    丢弃 session/model_change/bootstrap 等噪声，最后整体截断到 max_len。
    """
    if not text or text == "无 transcript 记录":
        return text or ""

    snippets: list[str] = []
    errors: list[str] = []  # 兜底：prompt-error 等异常事件
    for line in text.split("\n"):
        line = line.strip()
        if not line or not line.startswith("{"):
            continue
        try:
            ev = json.loads(line)
        except json.JSONDecodeError:
            continue
        etype = ev.get("type")
        # 异常事件兜底（运行失败、trajectory schema 无 message 时仍能给出线索）
        if etype == "custom" and "error" in str(ev.get("customType", "")).lower():
            data = ev.get("data", {})
            errors.append(f"[运行错误] {ev.get('customType')}: {json.dumps(data, ensure_ascii=False)[:300]}")
            continue
        if etype != "message":
            continue
        msg = ev.get("message", {}) or {}
        role = msg.get("role", "")
        content = msg.get("content", [])
        if not isinstance(content, list):
            continue
        for part in content:
            if not isinstance(part, dict):
                continue
            ptype = part.get("type")
            if ptype == "text" and part.get("text", "").strip():
                snippets.append(f"[{role}] {part['text'].strip()}")
            elif ptype == "thinking" and part.get("thinking", "").strip():
                snippets.append(f"[{role}:思考] {part['thinking'].strip()}")
            elif ptype == "toolCall":
                args = json.dumps(part.get("arguments", {}), ensure_ascii=False)
                snippets.append(f"[工具调用] {part.get('name', '')}({args[:500]})")
        if role == "toolResult":
            txts = [
                p.get("text", "") for p in content
                if isinstance(p, dict) and p.get("type") == "text"
            ]
            joined = " ".join(t for t in txts if t).strip()
            if joined:
                err = "(错误)" if msg.get("isError") else ""
                snippets.append(f"[工具结果{err}] {joined[:500]}")

    if not snippets and errors:
        snippets = errors
    result = "\n".join(snippets)
    if len(result) > max_len:
        result = result[:max_len] + "\n...（摘要已截断）"
    # 仍为空（如 trajectory schema 只有元数据事件）→ 回退原始文本截断
    return result or text[:max_len]



# --------------------------------------------------------------------------- #
# 样式辅助
# --------------------------------------------------------------------------- #
def style_header_row(ws, row: int, ncols: int):
    for col in range(1, ncols + 1):
        cell = ws.cell(row=row, column=col)
        cell.font = HEADER_FONT_WHITE
        cell.fill = HEADER_FILL
        cell.alignment = CENTER


def best_model_and_spread(scores: dict[str, float | None]) -> tuple[str, str]:
    """给定 {模型: 得分率}，返回 (最佳模型, 最高-最低分差字符串)。"""
    valid = {m: v for m, v in scores.items() if v is not None}
    if not valid:
        return "-", "-"
    best = max(valid, key=valid.get)
    spread = max(valid.values()) - min(valid.values())
    return best, f"{spread:.1f}%"


# --------------------------------------------------------------------------- #
# 用例顺序（按 manifest，回退首个模型的 tasks 顺序）
# --------------------------------------------------------------------------- #
def build_task_order(project_root: Path | None, models: list[ModelResult]) -> list[str]:
    """返回用例展示顺序（task_id 列表）。

    优先按 manifest.yaml 的 categories 顺序；无法读取时按首个模型的去重 tasks 顺序。
    只保留至少在一个模型中出现过的用例。
    """
    present: set[str] = set()
    for m in models:
        present.update(m.task_score.keys())

    ordered: list[str] = []
    if project_root is not None:
        mf = project_root / "tasks" / "manifest.yaml"
        if mf.exists():
            try:
                data = yaml.safe_load(mf.read_text(encoding="utf-8")) or {}
                cats = data.get("categories", {}) or {}
                for ids in cats.values():
                    for tid in ids or []:
                        if tid in present and tid not in ordered:
                            ordered.append(tid)
            except (OSError, yaml.YAMLError):
                ordered = []

    # 补齐 manifest 未覆盖但结果中出现的用例（按首个模型顺序）
    if models:
        for t in models[0].tasks:
            tid = t.get("task_id")
            if tid in present and tid not in ordered:
                ordered.append(tid)
    for tid in present:
        if tid not in ordered:
            ordered.append(tid)
    return ordered


# --------------------------------------------------------------------------- #
# Sheet 1: 总览
# --------------------------------------------------------------------------- #
def write_overview_sheet(wb, models, metas, task_order):
    ws = wb.active
    ws.title = "总览"

    # 各分类 / 场景 实际出现的（用于决定列）
    cat_to_tasks: dict[str, list[str]] = {}
    scene_to_tasks: dict[str, list[str]] = {}
    diff_to_tasks: dict[str, list[str]] = {}
    for tid in task_order:
        meta = metas.get(tid)
        if not meta:
            continue
        if meta.category_en:
            cat_to_tasks.setdefault(meta.category_en, []).append(tid)
        if meta.scene_en:
            scene_to_tasks.setdefault(meta.scene_en, []).append(tid)
        if meta.difficulty:
            diff_to_tasks.setdefault(meta.difficulty, []).append(tid)

    cats = [c for c in CATEGORY_ORDER if c in cat_to_tasks]
    scenes = [s for s in SCENE_ORDER if s in scene_to_tasks]
    # 难度固定展示 L1~L4 全列（无数据填 "-"），与参考产物一致
    diffs = list(DIFFICULTY_ORDER)

    headers = ["模型", "运行ID", "suite", "总分率", "用例数"]
    headers += [bilingual(CATEGORY_ZH.get(c, c), c) for c in cats]
    headers += [bilingual(DIFFICULTY_ZH.get(d, d), d) for d in diffs]
    headers += [bilingual(SCENE_ZH.get(s, s), s) for s in scenes]
    headers += ["总tokens", "总请求数", "总耗时(s)", "时间"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    for m in models:
        row = [
            m.model, m.run_id, m.suite,
            fmt_pct(m.total_pct), len(m.task_score),
        ]
        for c in cats:
            row.append(fmt_pct(avg_pct(m, cat_to_tasks[c])))
        for d in diffs:
            row.append(fmt_pct(avg_pct(m, diff_to_tasks.get(d, []))))
        for s in scenes:
            row.append(fmt_pct(avg_pct(m, scene_to_tasks[s])))
        eff = m.efficiency
        row += [
            eff.get("total_tokens", ""),
            eff.get("total_requests", ""),
            eff.get("total_execution_time_seconds", ""),
            fmt_timestamp(m.timestamp),
        ]
        ws.append(row)

    ws.column_dimensions["A"].width = 22
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 16
    ws.freeze_panes = "B2"


# --------------------------------------------------------------------------- #
# Sheet 2: 用例对比明细
# --------------------------------------------------------------------------- #
def write_case_compare_sheet(wb, models, metas, task_order):
    ws = wb.create_sheet("用例对比明细")

    headers = [
        "场景", "场景分组(S1~S8)", "子场景", "用例ID", "用例名称", "难度等级",
        "输入(Prompt)", "最优模型", "各模型最大分差",
    ]
    headers += [f"{m.model} 平均分" for m in models]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    for tid in task_order:
        meta = metas.get(tid) or TaskMeta()
        scene_cell = bilingual(SCENE_ZH.get(meta.scene_en, meta.scene_en), meta.scene_en)
        group = SCENE_GROUP.get(meta.scene_en, "")
        scores = {m.model: m.task_score.get(tid) for m in models}
        valid = {k: v for k, v in scores.items() if v is not None}
        if valid:
            best = max(valid, key=valid.get)
            spread = f"{max(valid.values()) - min(valid.values()):.3f}"
        else:
            best, spread = "-", "-"

        row = [
            scene_cell, group, meta.sub_scene_zh, tid, meta.name_zh, meta.difficulty,
            meta.prompt, best, spread,
        ]
        for m in models:
            v = m.task_score.get(tid)
            row.append(round(v, 3) if v is not None else "-")
        ws.append(row)
        r = ws.max_row
        ws.cell(row=r, column=7).alignment = WRAP_TOP  # 输入(Prompt)
    widths = {1: 28, 2: 16, 3: 22, 4: 30, 5: 22, 6: 10, 7: 50, 8: 18, 9: 16}
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    for col in range(10, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


# --------------------------------------------------------------------------- #
# Sheet 3: Agent能力对比
# --------------------------------------------------------------------------- #
def write_capability_sheet(wb, models, metas, task_order):
    ws = wb.create_sheet("Agent能力对比")

    # 收集实际出现的能力 -> 涉及用例
    cap_to_tasks: dict[str, list[str]] = {}
    for tid in task_order:
        meta = metas.get(tid)
        if not meta:
            continue
        for cap in meta.capabilities_en:
            cap_to_tasks.setdefault(cap, []).append(tid)
    caps = [c for c in CAPABILITY_ORDER if c in cap_to_tasks]
    caps += sorted(c for c in cap_to_tasks if c not in caps)

    headers = ["模型ID", "总分"]
    headers += [bilingual(CAPABILITY_ZH.get(c, c), c) for c in caps]
    headers += ["模型强项", "模型短板"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    for m in models:
        row = [m.model, fmt_pct(m.total_pct)]
        cap_pct: dict[str, float] = {}      # 全部有数据的能力得分率
        cap_pct_major: dict[str, float] = {}  # 仅涉及用例数 >= 阈值的能力（用于强项/短板）
        for c in caps:
            tids = [t for t in cap_to_tasks[c] if t in m.task_score]
            v = avg_pct(m, cap_to_tasks[c])
            cnt = len(tids)
            if v is not None:
                cap_pct[c] = v
                if cnt >= CAP_RANK_MIN_COUNT:
                    cap_pct_major[c] = v
                row.append(f"{v:.1f}% ({cnt})")
            else:
                row.append("-")
        # 强项 Top3 / 短板 Bottom3（仅在涉及用例数足够的能力中排名，避免小样本干扰）
        ranked = sorted(cap_pct_major.items(), key=lambda x: -x[1])
        top = ranked[:3]
        bottom = ranked[-3:][::-1]
        row.append("\n".join(
            f"{bilingual(CAPABILITY_ZH.get(c, c), c)}: {v:.1f}%" for c, v in top
        ))
        row.append("\n".join(
            f"{bilingual(CAPABILITY_ZH.get(c, c), c)}: {v:.1f}%" for c, v in bottom
        ))
        ws.append(row)
        r = ws.max_row
        ws.cell(row=r, column=len(headers) - 1).alignment = WRAP_TOP
        ws.cell(row=r, column=len(headers)).alignment = WRAP_TOP

    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 10
    for col in range(3, len(headers) - 1):
        ws.column_dimensions[get_column_letter(col)].width = 22
    ws.column_dimensions[get_column_letter(len(headers) - 1)].width = 40
    ws.column_dimensions[get_column_letter(len(headers))].width = 40
    ws.freeze_panes = "B2"


# --------------------------------------------------------------------------- #
# Sheet 4/5/6: 分类 / 场景 / 难度 对比（通用）
# --------------------------------------------------------------------------- #
def write_dimension_compare_sheet(
    wb, sheet_name, first_header, models, dim_keys, dim_to_tasks, label_fn
):
    """通用维度对比 Sheet。

    Args:
        first_header: 第一列表头（如"分类名称"）
        dim_keys: 维度键的有序列表
        dim_to_tasks: {维度键: [task_id]}
        label_fn: 维度键 -> 展示名
    """
    ws = wb.create_sheet(sheet_name)

    headers = [first_header, "用例数"]
    headers += [f"{m.model} 平均分" for m in models]
    headers += ["最佳模型", "最高-最低分差"]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    for key in dim_keys:
        tids = dim_to_tasks[key]
        row = [label_fn(key), len(tids)]
        scores = {}
        for m in models:
            v = avg_pct(m, tids)
            scores[m.model] = v
            row.append(fmt_pct(v))
        best, spread = best_model_and_spread(scores)
        row += [best, spread]
        ws.append(row)

    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 10
    for col in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = "B2"


def write_category_sheet(wb, models, metas, task_order):
    dim: dict[str, list[str]] = {}
    for tid in task_order:
        meta = metas.get(tid)
        if meta and meta.category_en:
            dim.setdefault(meta.category_en, []).append(tid)
    keys = [c for c in CATEGORY_ORDER if c in dim]
    keys += sorted(c for c in dim if c not in keys)
    write_dimension_compare_sheet(
        wb, "分类对比", "分类名称", models, keys, dim,
        lambda c: bilingual(CATEGORY_ZH.get(c, c), c),
    )


def write_scene_sheet(wb, models, metas, task_order):
    dim: dict[str, list[str]] = {}
    for tid in task_order:
        meta = metas.get(tid)
        if meta and meta.scene_en:
            dim.setdefault(meta.scene_en, []).append(tid)
    keys = [s for s in SCENE_ORDER if s in dim]
    keys += sorted(s for s in dim if s not in keys)

    def label(s):
        grp = SCENE_GROUP.get(s, "")
        name = bilingual(SCENE_ZH.get(s, s), s)
        return f"{grp} {name}".strip()

    write_dimension_compare_sheet(
        wb, "场景对比", "场景分组", models, keys, dim, label,
    )


def write_difficulty_sheet(wb, models, metas, task_order):
    dim: dict[str, list[str]] = {}
    for tid in task_order:
        meta = metas.get(tid)
        if meta and meta.difficulty:
            dim.setdefault(meta.difficulty, []).append(tid)
    keys = [d for d in DIFFICULTY_ORDER if d in dim]
    keys += sorted(d for d in dim if d not in keys)
    write_dimension_compare_sheet(
        wb, "难度等级对比", "难度等级", models, keys, dim,
        lambda d: bilingual(DIFFICULTY_ZH.get(d, d), d),
    )


# --------------------------------------------------------------------------- #
# Sheet 7: 模型分差矩阵
# --------------------------------------------------------------------------- #
def write_diff_matrix_sheet(wb, models):
    ws = wb.create_sheet("模型分差矩阵")

    headers = ["行模型 - 列模型"] + [m.model for m in models]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    for row_m in models:
        row = [row_m.model]
        for col_m in models:
            if row_m is col_m:
                row.append("0")
            else:
                diff = row_m.total_pct - col_m.total_pct
                row.append(f"{diff:+.1f}%")
        ws.append(row)

    ws.column_dimensions["A"].width = 24
    for col in range(2, len(headers) + 1):
        ws.column_dimensions[get_column_letter(col)].width = 18
    ws.freeze_panes = "B2"


# --------------------------------------------------------------------------- #
# Sheet 8..: 各模型评分详情
# --------------------------------------------------------------------------- #
def write_score_detail_sheet(wb, model, metas, task_order, analysis=None):
    """写入单个模型的评分详情。

    analysis: {"model::task_id": {"result_analysis": str, "root_cause": str}}
              用于回填「结果分析」「根因分析」两列；为 None 或缺项时留空。
    """
    analysis = analysis or {}
    # Sheet 名长度上限 31；模型名过长时截断
    name = f"评分详情_{model.model}"
    ws = wb.create_sheet(name[:31])

    headers = [
        "场景大类", "场景分组(S1~S8)", "子场景", "用例ID", "用例名称", "难度等级",
        "输入(Prompt)", "预期行为", "评分标准", "实际结果", "得分",
        "检查点得分明细", "失分点", "结果分析", "根因分析",
    ]
    ws.append(headers)
    style_header_row(ws, 1, len(headers))

    # task_id -> task dict（取去重后）
    task_by_id = {t["task_id"]: t for t in model.tasks if t.get("task_id")}

    for tid in task_order:
        meta = metas.get(tid) or TaskMeta()
        task = task_by_id.get(tid)
        if task is None:
            continue  # 该模型未跑此用例

        grading = task.get("grading", {}) or {}
        cat_cell = bilingual(CATEGORY_ZH.get(meta.category_en, meta.category_en), meta.category_en)
        group = SCENE_GROUP.get(meta.scene_en, "")
        transcript = read_transcript(model.path, model.run_id, tid)
        breakdown = format_breakdown(grading)
        lost_points = format_lost_points(grading)

        ana = analysis.get(f"{model.model}::{tid}", {})
        result_analysis = ana.get("result_analysis") or None
        root_cause = ana.get("root_cause") or None

        row = [
            cat_cell, group, meta.sub_scene_zh, tid, meta.name_zh, meta.difficulty,
            meta.prompt, meta.expected, meta.criteria, transcript,
            round(float(grading.get("mean", 0.0)), 3), breakdown,
            lost_points, result_analysis, root_cause,
        ]
        ws.append(row)
        r = ws.max_row
        for col in [7, 8, 9, 10, 12, 13, 14, 15]:  # 长文本列换行
            ws.cell(row=r, column=col).alignment = WRAP_TOP

    widths = {
        1: 22, 2: 16, 3: 22, 4: 30, 5: 22, 6: 10,
        7: 45, 8: 45, 9: 40, 10: 60, 11: 8, 12: 40, 13: 45, 14: 45, 15: 40,
    }
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = w
    ws.freeze_panes = "E2"


def fmt_timestamp(ts) -> str:
    try:
        return datetime.fromtimestamp(float(ts)).strftime("%Y-%m-%d %H:%M:%S")
    except (TypeError, ValueError, OSError):
        # 已是字符串时间戳则原样返回
        return str(ts) if ts else ""


# --------------------------------------------------------------------------- #
# 构建报告
# --------------------------------------------------------------------------- #
def build_report(models, metas, project_root, analysis=None) -> openpyxl.Workbook:
    task_order = build_task_order(project_root, models)

    wb = openpyxl.Workbook()
    write_overview_sheet(wb, models, metas, task_order)
    write_case_compare_sheet(wb, models, metas, task_order)
    write_capability_sheet(wb, models, metas, task_order)
    write_category_sheet(wb, models, metas, task_order)
    write_scene_sheet(wb, models, metas, task_order)
    write_difficulty_sheet(wb, models, metas, task_order)
    write_diff_matrix_sheet(wb, models)
    for m in models:
        write_score_detail_sheet(wb, m, metas, task_order, analysis)
    return wb


def build_analysis_input(models, metas, task_order) -> list[dict]:
    """构造「待 LLM 分析清单」：仅失分用例（mean < 1.0）。"""
    items = []
    for m in models:
        task_by_id = {t["task_id"]: t for t in m.tasks if t.get("task_id")}
        for tid in task_order:
            task = task_by_id.get(tid)
            if task is None:
                continue
            grading = task.get("grading", {}) or {}
            mean = float(grading.get("mean", 0.0))
            if mean >= 1.0:
                continue  # 满分用例不需分析
            meta = metas.get(tid) or TaskMeta()
            transcript_raw = read_transcript_raw(m.path, m.run_id, tid)
            items.append({
                "key": f"{m.model}::{tid}",
                "model": m.model,
                "task_id": tid,
                "task_name_zh": meta.name_zh,
                "category_zh": bilingual(
                    CATEGORY_ZH.get(meta.category_en, meta.category_en), meta.category_en
                ),
                "difficulty": meta.difficulty,
                "prompt_zh": meta.prompt,
                "expected_zh": meta.expected,
                "criteria_zh": meta.criteria,
                "score": round(mean, 3),
                "lost_points": format_lost_points(grading),
                "transcript_excerpt": summarize_transcript(transcript_raw),
                "result_analysis": "",
                "root_cause": "",
            })
    return items


def load_analysis(path: Path) -> dict:
    """读取已填写的分析 JSON，返回 {key: {result_analysis, root_cause}}。"""
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (json.JSONDecodeError, OSError) as e:
        print(f"警告：分析文件读取失败，将忽略: {path} -> {e}", file=sys.stderr)
        return {}
    out = {}
    for it in data.get("items", []):
        key = it.get("key")
        if key:
            out[key] = {
                "result_analysis": (it.get("result_analysis") or "").strip(),
                "root_cause": (it.get("root_cause") or "").strip(),
            }
    return out


# --------------------------------------------------------------------------- #
# 主流程
# --------------------------------------------------------------------------- #
def main():
    parser = argparse.ArgumentParser(
        description="PinchBench 评测报告汇总（中文整合版）"
    )
    parser.add_argument(
        "-d", "--dir", type=str, nargs="+", required=True,
        help="一个或多个结果目录，递归扫描 0xxx_<模型id>.json 文件",
    )
    parser.add_argument(
        "-o", "--output-dir", type=str, default=None,
        help="Excel 输出目录，默认 <结果根>/report-workspace/output（结果根从 -d 推断）",
    )
    parser.add_argument(
        "--analysis", type=str, default=None,
        help="已填写的分析 JSON 路径，用于回填「结果分析」「根因分析」两列",
    )
    args = parser.parse_args()

    # 收集结果文件
    all_files: list[Path] = []
    for d in args.dir:
        dp = Path(d)
        if not dp.is_dir():
            print(f"警告：不是目录，跳过: {dp}", file=sys.stderr)
            continue
        files = find_result_files(dp)
        if not files:
            print(f"警告：未找到符合命名规则的结果文件: {dp}", file=sys.stderr)
            continue
        print(f"在 {dp} 下找到 {len(files)} 个结果文件")
        all_files.extend(files)

    all_files = sorted(set(all_files))
    if not all_files:
        sys.exit("未找到任何符合命名规则（0xxx_<模型id>.json）的结果文件")

    # 加载模型结果
    models: list[ModelResult] = []
    for path in all_files:
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError) as e:
            print(f"跳过（解析失败）: {path} -> {e}", file=sys.stderr)
            continue
        models.append(ModelResult(path, data))

    if not models:
        sys.exit("没有可解析的结果文件")

    print(f"共加载 {len(models)} 个模型结果: {', '.join(m.model for m in models)}")

    # 加载用例中文元数据
    project_root = find_project_root(all_files[0].parent) or find_project_root(Path(__file__))
    if project_root is None:
        print("警告：未找到项目根，无法加载中文用例元数据，部分列将为空", file=sys.stderr)
        metas: dict[str, TaskMeta] = {}
    else:
        print(f"项目根: {project_root}")
        metas = load_task_meta(project_root)
        print(f"已加载 {len(metas)} 个用例的中文元数据")

    # 推断评测结果根目录：所有结果文件目录的共同父目录
    result_root = infer_result_root(all_files)
    print(f"评测结果根目录: {result_root}")

    # 读取分析回填（若提供）
    analysis = load_analysis(Path(args.analysis)) if args.analysis else None
    if analysis:
        print(f"已加载分析回填 {len(analysis)} 条")

    # 构建报告
    wb = build_report(models, metas, project_root, analysis)

    # 输出位置：Excel → <结果根>/report-workspace/output（-o 可覆盖）；JSON → <结果根>/report-workspace
    output_dir = Path(args.output_dir) if args.output_dir else result_root / "report-workspace" / "output"
    output_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = output_dir / f"report_{len(models)}models_{timestamp}.xlsx"
    wb.save(output_path)
    print(f"✓ 报告已生成: {output_path}")
    print(f"  - 模型数: {len(models)}")
    print(f"  - Sheet 数: {len(wb.sheetnames)} ({', '.join(wb.sheetnames)})")

    # 导出待分析清单（仅在未提供 --analysis 时，即首轮生成）
    if not args.analysis:
        task_order = build_task_order(project_root, models)
        items = build_analysis_input(models, metas, task_order)
        ws_dir = result_root / "report-workspace"
        ws_dir.mkdir(parents=True, exist_ok=True)
        ana_path = ws_dir / f"report_{len(models)}models_{timestamp}_analysis_input.json"
        ana_path.write_text(
            json.dumps(
                {"report_file": output_path.name, "items": items},
                ensure_ascii=False, indent=2,
            ),
            encoding="utf-8",
        )
        print(f"✓ 待分析清单已导出: {ana_path}")
        print(f"  - 失分用例数: {len(items)}")
        print("  - 填写 result_analysis/root_cause 后，用 --analysis 重跑以回填")


if __name__ == "__main__":
    main()
