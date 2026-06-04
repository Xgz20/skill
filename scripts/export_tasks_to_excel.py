#!/usr/bin/env python3
"""
PinchBench 评测任务导出脚本

遍历所有评测任务，提取关键信息并导出到 Excel 文件。
"""

import argparse
import logging
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    sys.exit("需要 openpyxl 库，请执行: pip install openpyxl")

# 添加 scripts 目录到 Python 路径以导入 lib_tasks
sys.path.insert(0, str(Path(__file__).parent))

from lib_tasks import TaskLoader


logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s"
)
logger = logging.getLogger(__name__)


def format_workspace_files(workspace_files: list) -> str:
    """格式化工作区文件列表为可读字符串。"""
    if not workspace_files:
        return "无"

    files = []
    for item in workspace_files:
        if isinstance(item, dict):
            # 格式: {source: ..., dest: ...} 或 {path: ..., content: ...}
            if "source" in item:
                files.append(f"{item.get('dest', item['source'])}")
            elif "path" in item:
                files.append(item["path"])
        elif isinstance(item, str):
            files.append(item)

    return "\n".join(files) if files else "无"


def format_grading_criteria(criteria: list) -> str:
    """格式化评分标准列表为带编号的字符串。"""
    if not criteria:
        return "无"
    return "\n".join(f"{i+1}. {item}" for i, item in enumerate(criteria))


def truncate_text(text: str, max_length: int = 500) -> str:
    """截断过长的文本。"""
    if not text:
        return ""
    if len(text) <= max_length:
        return text
    return text[:max_length] + "..."


def write_tasks_sheet(wb: openpyxl.Workbook, tasks: list, core_task_ids: set):
    """写入任务详情工作表。"""
    ws = wb.active
    ws.title = "任务详情"

    # 定义表头样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

    # 定义表头
    headers = [
        "序号", "分类", "任务ID", "任务名称", "评分类型",
        "超时(秒)", "核心任务", "多轮对话", "输入文件",
        "任务提示", "预期行为", "评分标准"
    ]

    ws.append(headers)

    # 设置表头样式
    for col_idx, _ in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 写入任务数据
    for idx, task in enumerate(tasks, 1):
        is_core = "是" if task.task_id in core_task_ids else "否"
        is_multi_session = "是" if task.frontmatter.get("multi_session", False) else "否"

        row_data = [
            idx,
            task.category,
            task.task_id,
            task.name,
            task.grading_type,
            task.timeout_seconds,
            is_core,
            is_multi_session,
            format_workspace_files(task.workspace_files),
            truncate_text(task.prompt, 1000),
            truncate_text(task.expected_behavior, 1000),
            format_grading_criteria(task.grading_criteria),
        ]

        ws.append(row_data)

        # 设置文本换行（对于长文本列）
        row_num = idx + 1
        for col_idx in [9, 10, 11, 12]:  # 输入文件、提示、预期行为、评分标准
            cell = ws.cell(row=row_num, column=col_idx)
            cell.alignment = Alignment(wrap_text=True, vertical="top")

    # 设置列宽
    column_widths = {
        1: 6,   # 序号
        2: 15,  # 分类
        3: 35,  # 任务ID
        4: 30,  # 任务名称
        5: 12,  # 评分类型
        6: 10,  # 超时
        7: 10,  # 核心任务
        8: 10,  # 多轮对话
        9: 25,  # 输入文件
        10: 50, # 任务提示
        11: 50, # 预期行为
        12: 50, # 评分标准
    }

    for col_idx, width in column_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    # 冻结首行
    ws.freeze_panes = "A2"

    # 启用自动筛选
    ws.auto_filter.ref = ws.dimensions

    logger.info(f"已写入 {len(tasks)} 个任务到任务详情工作表")


def write_statistics_sheet(wb: openpyxl.Workbook, tasks: list, categories: list):
    """写入统计信息工作表。"""
    ws = wb.create_sheet("统计信息")

    # 定义样式
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    title_font = Font(bold=True, size=12)

    # 1. 总体统计
    ws.append(["总体统计"])
    ws["A1"].font = title_font
    ws.append([])

    ws.append(["指标", "数值"])
    ws["A3"].font = header_font
    ws["A3"].fill = header_fill
    ws["B3"].font = header_font
    ws["B3"].fill = header_fill

    ws.append(["任务总数", len(tasks)])
    ws.append(["分类数量", len(categories)])

    # 统计评分类型
    grading_types = {}
    multi_session_count = 0
    for task in tasks:
        grading_types[task.grading_type] = grading_types.get(task.grading_type, 0) + 1
        if task.frontmatter.get("multi_session", False):
            multi_session_count += 1

    ws.append(["多轮对话任务", multi_session_count])
    ws.append([])

    # 2. 按评分类型统计
    ws.append(["按评分类型统计"])
    ws[f"A{ws.max_row}"].font = title_font
    ws.append([])

    ws.append(["评分类型", "任务数", "占比"])
    row_num = ws.max_row
    for col_idx in range(1, 4):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    for grading_type, count in sorted(grading_types.items()):
        percentage = f"{count / len(tasks) * 100:.1f}%"
        ws.append([grading_type, count, percentage])

    ws.append([])

    # 3. 按分类统计
    ws.append(["按分类统计"])
    ws[f"A{ws.max_row}"].font = title_font
    ws.append([])

    ws.append(["分类", "任务数", "占比"])
    row_num = ws.max_row
    for col_idx in range(1, 4):
        cell = ws.cell(row=row_num, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill

    # 统计每个分类的任务数
    category_counts = {}
    for task in tasks:
        category_counts[task.category] = category_counts.get(task.category, 0) + 1

    # 按分类顺序输出（如果有的话），否则按字母顺序
    if categories:
        sorted_categories = [(cat, category_counts.get(cat, 0)) for cat in categories]
    else:
        sorted_categories = sorted(category_counts.items())

    for category, count in sorted_categories:
        if count > 0:  # 只显示有任务的分类
            percentage = f"{count / len(tasks) * 100:.1f}%"
            ws.append([category, count, percentage])

    # 设置列宽
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 12

    logger.info("已写入统计信息工作表")


def main():
    parser = argparse.ArgumentParser(
        description="导出 PinchBench 评测任务到 Excel 文件"
    )
    parser.add_argument(
        "--output-dir",
        type=str,
        default=None,
        help="输出目录，默认为 docs/cases/"
    )
    parser.add_argument(
        "--core-only",
        action="store_true",
        help="仅导出核心任务"
    )
    parser.add_argument(
        "--category",
        type=str,
        default=None,
        help="仅导出指定分类的任务"
    )
    parser.add_argument(
        "--grading-type",
        type=str,
        default=None,
        choices=["automated", "llm_judge", "hybrid"],
        help="仅导出指定评分类型的任务 (automated/llm_judge/hybrid)"
    )

    args = parser.parse_args()

    # 确定项目根目录和任务目录
    script_dir = Path(__file__).resolve().parent
    project_root = script_dir.parent
    tasks_dir = project_root / "tasks"

    if not tasks_dir.exists():
        logger.error(f"任务目录不存在: {tasks_dir}")
        sys.exit(1)

    # 加载所有任务
    logger.info(f"从 {tasks_dir} 加载任务...")
    loader = TaskLoader(tasks_dir)
    all_tasks = loader.load_all_tasks()

    if not all_tasks:
        logger.error("未找到任何任务")
        sys.exit(1)

    logger.info(f"成功加载 {len(all_tasks)} 个任务")

    # 过滤任务
    tasks = all_tasks
    core_task_ids = set(loader.core_tasks)

    if args.core_only:
        tasks = [t for t in tasks if t.task_id in core_task_ids]
        logger.info(f"过滤后剩余 {len(tasks)} 个核心任务")

    if args.category:
        tasks = [t for t in tasks if t.category == args.category]
        logger.info(f"过滤后剩余 {len(tasks)} 个 {args.category} 分类的任务")

    if args.grading_type:
        tasks = [t for t in tasks if t.grading_type == args.grading_type]
        logger.info(f"过滤后剩余 {len(tasks)} 个 {args.grading_type} 评分类型的任务")

    if not tasks:
        logger.error("过滤后没有任务可导出")
        sys.exit(1)

    # 确定输出目录
    if args.output_dir:
        output_dir = Path(args.output_dir)
    else:
        output_dir = project_root / "docs" / "cases"

    output_dir.mkdir(parents=True, exist_ok=True)
    logger.info(f"输出目录: {output_dir}")

    # 生成文件名：根据过滤参数构造
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    name_parts = ["task_details"]
    if args.core_only:
        name_parts.append("core")
    if args.category:
        name_parts.append(args.category)
    if args.grading_type:
        name_parts.append(args.grading_type)
    if not args.core_only and not args.category and not args.grading_type:
        name_parts = ["all_task_details"]
    name_parts.append(timestamp)
    output_filename = "_".join(name_parts) + ".xlsx"
    output_path = output_dir / output_filename

    # 创建 Excel 工作簿
    logger.info("创建 Excel 工作簿...")
    wb = openpyxl.Workbook()

    # 写入任务详情
    write_tasks_sheet(wb, tasks, core_task_ids)

    # 写入统计信息
    write_statistics_sheet(wb, tasks, loader.categories)

    # 保存文件
    wb.save(output_path)
    logger.info(f"✓ 导出完成: {output_path}")
    logger.info(f"  - 任务总数: {len(tasks)}")
    logger.info(f"  - 核心任务: {len([t for t in tasks if t.task_id in core_task_ids])}")
    logger.info(f"  - 分类数量: {len(set(t.category for t in tasks))}")


if __name__ == "__main__":
    main()
